"""Generates sample test PDF and DOCX documents using ReportLab and python-docx."""

import io

import pypdf
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def create_sample_pdf(text: str = "John Doe\nSoftware Engineer\nSkills: Python, React, PostgreSQL\nExperience: Jan 2021 - Present") -> bytes:
    """Generate a clean single-page PDF with text."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    y = 750
    for line in text.splitlines():
        c.drawString(100, y, line)
        y -= 20
    c.showPage()
    c.save()
    return buffer.getvalue()


def create_multipage_pdf(pages_text: list[str]) -> bytes:
    """Generate a multi-page PDF."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    for p_text in pages_text:
        y = 750
        for line in p_text.splitlines():
            c.drawString(100, y, line)
            y -= 20
        c.showPage()
    c.save()
    return buffer.getvalue()


def create_scanned_image_pdf() -> bytes:
    """Generate an image-only PDF with no text layer."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    # Draw a rectangle representing a scanned image, but no text string
    c.rect(50, 50, 500, 700, fill=1)
    c.showPage()
    c.save()
    return buffer.getvalue()


def create_encrypted_pdf(text: str = "Confidential resume") -> bytes:
    """Generate a password-protected PDF."""
    plain_pdf = create_sample_pdf(text)
    reader = pypdf.PdfReader(io.BytesIO(plain_pdf))
    writer = pypdf.PdfWriter()
    for p in reader.pages:
        writer.add_page(p)
    writer.encrypt("secret_pass")
    out_buffer = io.BytesIO()
    writer.write(out_buffer)
    return out_buffer.getvalue()


def create_sample_docx(text: str = "Jane Smith\nSenior Python Developer\nSkills: Python, FastAPI, Docker, Kubernetes\nWork Experience: 2019 - 2023 at TechCorp") -> bytes:
    """Generate a valid DOCX document."""
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)

    # Add a table
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Skill"
    t.cell(0, 1).text = "Level"
    t.cell(1, 0).text = "Python"
    t.cell(1, 1).text = "Expert"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
